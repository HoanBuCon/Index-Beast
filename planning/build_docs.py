"""Sinh docs/03_SPRINT_PLAN.md từ sprint_plan.json và xuất các tài liệu docs/*.md sang .docx.

Cách dùng:
    pip install -r planning/requirements.txt
    python planning/build_docs.py
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
PLAN_PATH = ROOT / "planning" / "sprint_plan.json"
DOCS_DIR = ROOT / "docs"
EXPORT_DIR = DOCS_DIR / "export"

PRIORITIES = ("Must", "Should", "Could")


def load_plan() -> dict:
    return json.loads(PLAN_PATH.read_text(encoding="utf-8"))


def sprint_dates(plan: dict, number: int) -> tuple[date, date]:
    start = date.fromisoformat(plan["start_date"]) + timedelta(days=(number - 1) * plan["sprint_length_days"])
    return start, start + timedelta(days=plan["sprint_length_days"] - 1)


def priority_of(card: dict) -> str:
    return next((p for p in PRIORITIES if p in card.get("labels", [])), "-")


def fmt(d: date) -> str:
    return d.strftime("%d/%m/%Y")


# ---------------------------------------------------------------------------
# Sprint plan markdown
# ---------------------------------------------------------------------------

def build_sprint_plan_md(plan: dict) -> str:
    members = [m["id"] for m in plan["members"]]
    out: list[str] = []
    w = out.append

    first_start, _ = sprint_dates(plan, 1)
    _, last_end = sprint_dates(plan, len(plan["sprints"]))

    w(f"# {plan['project']} — Kế hoạch 6 Sprint")
    w("")
    w("| Thuộc tính | Giá trị |")
    w("|---|---|")
    w(f"| Thời gian | {fmt(first_start)} → {fmt(last_end)} ({len(plan['sprints'])} sprint × {plan['sprint_length_days']} ngày) |")
    w(f"| Quy mô | {len(members)} thành viên |")
    w("| Nguồn dữ liệu kế hoạch | planning/sprint_plan.json (sửa file này rồi chạy lại build_docs.py / trello_import.py) |")
    w("| Trello | https://trello.com/b/mZEBMPGV/elden-ring |")
    w("")
    w("> Tài liệu này được sinh tự động từ `planning/sprint_plan.json`. Đừng sửa tay — sửa JSON rồi chạy lại script.")
    w("")

    w("## 1. Phân vai")
    w("")
    w("| Mã | Vai trò | Trọng tâm |")
    w("|---|---|---|")
    for m in plan["members"]:
        w(f"| {m['id']} | {m['role']} | {m['focus']} |")
    w("")
    w("Nguyên tắc: mỗi thẻ có **một** người chịu trách nhiệm chính; thẻ `ALL` là hoạt động chung của team. "
      "Người chịu trách nhiệm có thể nhờ người khác hỗ trợ nhưng là người kéo thẻ sang Hoàn thành.")
    w("")

    w("## 2. Lộ trình tổng quan")
    w("")
    w("| Sprint | Thời gian | Chủ đề | Mục tiêu |")
    w("|---|---|---|---|")
    for s in plan["sprints"]:
        a, b = sprint_dates(plan, s["number"])
        w(f"| S{s['number']} | {fmt(a)} – {fmt(b)} | **{s['name']}** | {s['goal']} |")
    w("")
    w("Các mốc quan trọng: **cuối S1** chốt kỹ thuật (ADR) · **cuối S3** MVP nội bộ chạy end-to-end · "
      "**cuối S4** đạt mục tiêu chất lượng · **cuối S5** production · **S6** UAT, release v1.0 và buffer.")
    w("")

    w("## 3. Phân bổ story point theo sprint")
    w("")
    load: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for s in plan["sprints"]:
        for c in s["cards"]:
            load[s["number"]][c["owner"]] += c["points"]
    header = ["Sprint", *members, "ALL", "Tổng"]
    w("| " + " | ".join(header) + " |")
    w("|" + "---|" * len(header))
    for s in plan["sprints"]:
        row = load[s["number"]]
        cells = [f"S{s['number']}", *(str(row[m]) for m in members), str(row["ALL"]), str(sum(row.values()))]
        w("| " + " | ".join(cells) + " |")
    totals = [sum(load[s["number"]][m] for s in plan["sprints"]) for m in [*members, "ALL"]]
    w("| **Tổng** | " + " | ".join(f"**{t}**" for t in totals) + f" | **{sum(totals)}** |")
    w("")
    w("Thang điểm Fibonacci (1, 2, 3, 5, 8): 1 điểm ≈ nửa ngày làm việc tập trung. "
      "Mỗi người ~10–14 điểm/sprint, chừa ~20% cho review code, họp và việc phát sinh. "
      "Sau Sprint 1, hiệu chỉnh theo velocity thực tế.")
    w("")

    w("## 4. Cấu trúc bảng Trello")
    w("")
    w("| Cột (List) | Mục đích |")
    w("|---|---|")
    purposes = {
        "📌 Thông tin dự án": "Link tài liệu, Definition of Done, quy ước Git, phân vai",
        "🧊 Backlog": "Ý tưởng / yêu cầu Could-have / CHANGE-REQ chưa xếp sprint",
        "🔨 Đang làm": "Thẻ đang thực hiện (mỗi người tối đa 2 thẻ cùng lúc)",
        "👀 Review / QA": "Đã mở PR, chờ review hoặc kiểm thử",
        "✅ Hoàn thành": "Đạt Definition of Done",
    }
    for name in plan["lists"]:
        w(f"| {name} | {purposes.get(name, 'Thẻ đã lên kế hoạch cho sprint này; kéo sang Đang làm khi bắt đầu')} |")
    w("")
    w("Nhãn: " + ", ".join(f"`{l['name']}`" for l in plan["labels"]) + ".")
    w("")
    w("Mỗi thẻ có: tên `[S#][TV#] Tên việc`, mô tả (mục tiêu, yêu cầu SRS liên quan, story point), checklist các việc con, "
      "ngày bắt đầu/hạn = đầu/cuối sprint, nhãn người phụ trách + mảng + ưu tiên.")
    w("")

    for s in plan["sprints"]:
        a, b = sprint_dates(plan, s["number"])
        w(f"## Sprint {s['number']} — {s['name']} ({fmt(a)} – {fmt(b)})")
        w("")
        w(f"**Mục tiêu:** {s['goal']}")
        w("")
        w("**Sản phẩm bàn giao cuối sprint:**")
        w("")
        for d in s["deliverables"]:
            w(f"- {d}")
        w("")
        w("| # | Thẻ | Phụ trách | Điểm | Ưu tiên | Yêu cầu SRS |")
        w("|---|---|---|---|---|---|")
        for i, c in enumerate(s["cards"], 1):
            srs = ", ".join(c.get("srs", [])) or "-"
            w(f"| {i} | {c['title']} | {c['owner']} | {c['points']} | {priority_of(c)} | {srs} |")
        w("")
        w(f"### Chi tiết công việc Sprint {s['number']}")
        w("")
        for i, c in enumerate(s["cards"], 1):
            w(f"#### {s['number']}.{i} [{c['owner']}] {c['title']}")
            w("")
            if c.get("desc"):
                w(c["desc"])
                w("")
            for item in c["checklist"]:
                w(f"- [ ] {item}")
            w("")

    w("## Backlog (chưa xếp sprint)")
    w("")
    w("| Thẻ | Phụ trách đề xuất | Điểm | Ghi chú |")
    w("|---|---|---|---|")
    for c in plan["backlog_cards"]:
        w(f"| {c['title']} | {c['owner']} | {c['points']} | {c['desc'].splitlines()[0]} |")
    w("")

    w("## Nghi thức Scrum rút gọn (sprint 1 tuần)")
    w("")
    w("| Hoạt động | Thời điểm | Thời lượng | Nội dung |")
    w("|---|---|---|---|")
    w("| Sprint Planning | Thứ Hai đầu sprint | 45 phút | Chọn thẻ, xác nhận người phụ trách & điểm |")
    w("| Daily (async) | Hằng ngày trên Discord | 5 phút/người | Hôm qua / Hôm nay / Vướng mắc |")
    w("| Sync giữa sprint | Thứ Năm | 20 phút | Gỡ vướng, điều chỉnh phạm vi |")
    w("| Sprint Review | Chủ nhật | 30 phút | Demo, chạy eval, cập nhật SRS |")
    w("| Retrospective | Chủ nhật (sau Review) | 20 phút | Giữ / Bỏ / Thử |")
    w("")

    w("## Rủi ro kế hoạch & phương án dự phòng")
    w("")
    w("| Rủi ro | Dấu hiệu | Phương án |")
    w("|---|---|---|")
    w("| Crawl/cleaning lâu hơn dự kiến | S2 chưa index xong toàn bộ | S3 chạy trên tập con đã index; TV3 hỗ trợ TV2 |")
    w("| Rerank CPU quá chậm | p95 > 1 s ở S3 | Đổi MiniLM, giảm top-N, cache; quyết định ở Review S3 |")
    w("| Chất lượng chưa đạt ở S4 | Recall@5 < 0.8 | Ưu tiên dữ liệu & retrieval, hoãn các thẻ Could |")
    w("| Thành viên bận/ốm | Velocity giảm > 30% | Cắt thẻ Should/Could; S6 là buffer |")
    w("| Yêu cầu thay đổi nhiều | Nhiều thẻ CHANGE-REQ | Xử lý theo SRS §9, chỉ nhận vào sprint kế tiếp |")
    w("")
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------------------
# Markdown -> docx (đủ cho các tài liệu trong docs/)
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text: str) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], text)
        for r in cell.paragraphs[0].runs:
            r.bold = True
        shade(cell, "D9C48A")
    for row in body:
        cells = table.add_row().cells
        for i in range(len(header)):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], row[i] if i < len(row) else "")
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph()


def md_to_docx(md: str, target: Path, landscape: bool = False) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if landscape:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?\s*:?-{3,}.*", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            level = len(heading.group(1))
            h = doc.add_heading(level=0 if level == 1 else level - 1)
            add_inline(h, heading.group(2))
            if level == 1:
                for r in h.runs:
                    r.font.color.rgb = RGBColor(0x8A, 0x6D, 0x1E)
        elif stripped.startswith("- [ ] "):
            add_inline(doc.add_paragraph(style="List Bullet"), "☐ " + stripped[6:])
        elif stripped.startswith(("- ", "* ")):
            add_inline(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, stripped.lstrip("> "))
        elif stripped:
            add_inline(doc.add_paragraph(), stripped)
        i += 1

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def main() -> None:
    plan = load_plan()
    plan_md = DOCS_DIR / "03_SPRINT_PLAN.md"
    plan_md.write_text(build_sprint_plan_md(plan), encoding="utf-8")
    print(f"wrote {plan_md.relative_to(ROOT)}")

    for md_path in sorted(DOCS_DIR.glob("0*_*.md")):
        target = EXPORT_DIR / (md_path.stem + ".docx")
        md_to_docx(md_path.read_text(encoding="utf-8"), target, landscape=md_path.stem.startswith("03"))
        print(f"wrote {target.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
