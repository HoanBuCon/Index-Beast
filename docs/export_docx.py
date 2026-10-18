#!/usr/bin/env python3
"""Chuyển đổi tất cả file Markdown (.md) trong thư mục docs/ thành file Word (.docx) trong docs/export/.

Cách dùng:
    pip install -r docs/requirements.txt
    python docs/export_docx.py
    # Hoặc khi đang ở trong thư mục docs/:
    python export_docx.py

Tùy chọn:
    python docs/export_docx.py [file.md]       # Chỉ chuyển đổi file cụ thể
    python docs/export_docx.py --landscape     # Ép tất cả sang khổ ngang
    python docs/export_docx.py --no-recurse    # Chỉ chuyển đổi file .md ở cấp docs/ (không đệ quy)
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

# Đảm bảo in ký tự tiếng Việt an toàn trên Windows console
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print(
        "Lỗi: Thư viện 'python-docx' chưa được cài đặt.\n"
        "Vui lòng chạy: pip install -r docs/requirements.txt",
        file=sys.stderr,
    )
    sys.exit(1)

DOCS_DIR = Path(__file__).resolve().parent
EXPORT_DIR = DOCS_DIR / "export"

COLOR_GOLD = RGBColor(0x8A, 0x6D, 0x1E)
COLOR_TEXT = RGBColor(0x24, 0x29, 0x2F)
COLOR_LINK = RGBColor(0x09, 0x69, 0xDA)
HEX_HEADER_SHADE = "D9C48A"
HEX_CODE_BG = "F6F8FA"

TOKEN_RE = re.compile(
    r"(?P<bold>\*\*[^*]+\*\*)"
    r"|(?P<code>`[^`]+`)"
    r"|(?P<link>\[[^\]]+\]\([^)]+\))"
    r"|(?P<italic>(?<!\*)\*[^*]+\*(?!\*))"
)


def add_inline(paragraph, text: str) -> None:
    """Xử lý định dạng inline: bold, code, link, italic."""
    last_end = 0
    for match in TOKEN_RE.finditer(text):
        start, end = match.span()
        if start > last_end:
            paragraph.add_run(text[last_end:start])

        if match.group("bold"):
            val = match.group("bold")[2:-2]
            paragraph.add_run(val).bold = True
        elif match.group("code"):
            val = match.group("code")[1:-1]
            run = paragraph.add_run(val)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif match.group("link"):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", match.group("link"))
            if m:
                label = m.group(1)
                run = paragraph.add_run(label)
                run.font.color.rgb = COLOR_LINK
                run.underline = True
            else:
                paragraph.add_run(match.group("link"))
        elif match.group("italic"):
            val = match.group("italic")[1:-1]
            paragraph.add_run(val).italic = True

        last_end = end

    if last_end < len(text):
        paragraph.add_run(text[last_end:])


def shade_cell(cell, hex_color: str) -> None:
    """Đổ màu nền (shading) cho một cell trong bảng."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def split_table_row(line: str) -> list[str]:
    """Tách hàng bảng Markdown thành danh sách chuỗi."""
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]], landscape: bool = False) -> None:
    """Tạo bảng Word từ các hàng markdown đã parse."""
    if not rows:
        return

    header = rows[0]
    body = rows[1:]
    col_count = len(header)
    table = doc.add_table(rows=1, cols=col_count)
    table.style = "Table Grid"

    # Header
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, text)
        for r in p.runs:
            r.bold = True
        shade_cell(cell, HEX_HEADER_SHADE)

    # Body
    for row in body:
        row_cells = table.add_row().cells
        for col_idx in range(col_count):
            cell = row_cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            text_val = row[col_idx] if col_idx < len(row) else ""
            add_inline(p, text_val)

    # Font size & cell formatting
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)

    # Phân bổ độ rộng cột hợp lý
    avail_width = Inches(9.0 if landscape else 6.5)
    col_width = int(avail_width // col_count)
    for col in table.columns:
        col.width = col_width
        for cell in col.cells:
            cell.width = col_width

    doc.add_paragraph()


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    """Tạo block code có nền xám nhạt và font Consolas."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    shade_cell(cell, HEX_CODE_BG)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_TEXT
    doc.add_paragraph()


def should_use_landscape(md_text: str, file_path: Path) -> bool:
    """Tự động phát hiện nếu tài liệu cần khổ ngang (bảng nhiều cột hoặc sprint plan)."""
    if file_path.stem.startswith("03") or "sprint" in file_path.stem.lower():
        return True
    for line in md_text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and not re.fullmatch(r"\|?\s*:?-{3,}.*", stripped):
            cols = split_table_row(stripped)
            if len(cols) >= 6:
                return True
    return False


def md_to_docx(md_text: str, target: Path, landscape: bool = False) -> None:
    """Chuyển đổi nội dung Markdown thành file .docx."""
    doc = Document()

    # Style mặc định
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = COLOR_TEXT

    # Khổ giấy & lề
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)

    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block ```
        if stripped.startswith("```"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        # Bảng | ... |
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row_line = lines[i].strip()
                if not re.fullmatch(r"\|?\s*:?-{3,}.*", row_line):
                    rows.append(split_table_row(row_line))
                i += 1
            add_table(doc, rows, landscape=landscape)
            continue

        # Tiêu đề #, ##, ###, ####
        heading_match = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading_match:
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2)
            h = doc.add_heading(level=0 if level == 1 else level - 1)
            add_inline(h, heading_text)
            if level == 1:
                for r in h.runs:
                    r.font.color.rgb = COLOR_GOLD
            i += 1
            continue

        # Checkbox: - [ ] hoặc - [x]
        if stripped.startswith(("- [ ] ", "* [ ] ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, "☐ " + stripped[6:])
        elif stripped.startswith(("- [x] ", "* [x] ", "- [X] ", "* [X] ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, "☑ " + stripped[6:])
        # Bullet list
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
        # Numbered list
        elif re.match(r"^\d+\.\s", stripped):
            clean_text = re.sub(r"^\d+\.\s", "", stripped)
            p = doc.add_paragraph(style="List Number")
            add_inline(p, clean_text)
        # Blockquote >
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, stripped.lstrip("> "))
            for r in p.runs:
                r.font.italic = True
        # Đoạn văn bản thông thường
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, stripped)

        i += 1

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def collect_markdown_files(recurse: bool = True) -> list[Path]:
    """Tìm tất cả file markdown trong docs/, loại trừ thư mục export/."""
    pattern = "**/*.md" if recurse else "*.md"
    files: list[Path] = []
    for p in sorted(DOCS_DIR.glob(pattern)):
        # Bỏ qua nếu nằm trong export
        if EXPORT_DIR in p.parents or p.parent == EXPORT_DIR:
            continue
        files.append(p)
    return files


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Chuyển đổi file Markdown (.md) trong thư mục docs sang .docx trong docs/export/"
    )
    parser.add_argument(
        "file",
        nargs="?",
        default=None,
        help="Đường dẫn file .md cụ thể cần chuyển đổi (mặc định: tất cả file trong docs/)",
    )
    parser.add_argument(
        "--landscape",
        action="store_true",
        help="Bắt buộc xuất theo khổ ngang (mặc định: tự động phát hiện)",
    )
    parser.add_argument(
        "--no-recurse",
        action="store_true",
        help="Không quét đệ quy các thư mục con trong docs/",
    )
    args = parser.parse_args()

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    if args.file:
        target_path = Path(args.file).resolve()
        if not target_path.exists():
            print(f"Lỗi: Không tìm thấy file {target_path}", file=sys.stderr)
            sys.exit(1)
        files = [target_path]
    else:
        files = collect_markdown_files(recurse=not args.no_recurse)

    if not files:
        print("Không tìm thấy file markdown nào để chuyển đổi.")
        return

    print(f"Tìm thấy {len(files)} file markdown. Bắt đầu xuất sang {EXPORT_DIR.relative_to(DOCS_DIR.parent)}...\n")

    success_count = 0
    for md_file in files:
        # Đường dẫn tương đối so với docs/
        try:
            rel = md_file.relative_to(DOCS_DIR)
        except ValueError:
            rel = Path(md_file.name)

        out_path = EXPORT_DIR / rel.with_suffix(".docx")

        md_text = md_file.read_text(encoding="utf-8")
        is_landscape = args.landscape or should_use_landscape(md_text, md_file)

        orientation_str = "khổ ngang" if is_landscape else "khổ dọc"
        print(f"  • {rel.as_posix()} -> {out_path.name} ({orientation_str})...", end=" ")

        try:
            md_to_docx(md_text, out_path, landscape=is_landscape)
            print("OK")
            success_count += 1
        except Exception as e:
            print(f"LỖI ({e})")

    print(f"\nHoàn thành: {success_count}/{len(files)} file đã được chuyển đổi thành công.")


if __name__ == "__main__":
    main()
